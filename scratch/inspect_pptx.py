from pptx import Presentation

prs = Presentation(r'H:\Downloads\project  review  phase-II-P2  ppts Format.pptx')
print(f'Slide count: {len(prs.slides)}')
for i, slide in enumerate(prs.slides, 1):
    print(f'\n================ Slide {i} ================')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = p.text.strip()
                if text:
                    print(f'[{shape.name}] {text}')
