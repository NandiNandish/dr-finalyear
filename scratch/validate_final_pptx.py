from pptx import Presentation

prs = Presentation(r'H:\Downloads\project_review_phase_II_P2_colorful_presentation.pptx')
print(f'Total Slides in Final PPTX: {len(prs.slides)}')
for idx, slide in enumerate(prs.slides, 1):
    title = slide.shapes.title.text.strip() if slide.shapes.title else 'NO TITLE'
    pics = [s.name for s in slide.shapes if s.shape_type == 13]
    shapes_cnt = len(slide.shapes)
    print(f'Slide {idx:02d}: "{title}" (Total shapes: {shapes_cnt}, Pictures: {len(pics)})')
