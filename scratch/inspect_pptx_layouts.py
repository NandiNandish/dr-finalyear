from pptx import Presentation

prs = Presentation(r'H:\Downloads\project  review  phase-II-P2  ppts Format.pptx')
for i, slide in enumerate(prs.slides, 1):
    print(f'Slide {i} layout: {slide.slide_layout.name}, shapes count: {len(slide.shapes)}')
    for shape in slide.shapes:
        ph = shape.placeholder_format.type if shape.is_placeholder else "None"
        print(f'   shape: {shape.name}, type: {shape.shape_type}, placeholder: {ph}')
